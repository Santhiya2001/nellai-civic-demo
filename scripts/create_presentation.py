"""Generate an editable ME project presentation with native PowerPoint diagrams."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "demo-ready"
OUT.mkdir(parents=True, exist_ok=True)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "Nellai Green & Civic — ME Project Demonstration"
prs.core_properties.subject = "Modular civic and environmental management platform"
prs.core_properties.author = "Nellai Green & Civic Project"
GREEN = "214C36"
DARK = "173E2B"
CREAM = "F6F7F0"
SAGE = "E6EDDB"
MUTED = "6F8068"
INK = "283F30"
ORANGE = "D98551"
WHITE = "FFFFFF"
BLUE = "DCEDEC"
GOLD = "F3E9D3"


def color(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill=WHITE, border=None, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if border:
        shape.line.color.rgb = color(border)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.12
    return shape


def txt(slide, value, x, y, w, h, size=18, fill=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.01)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(fill)
        p.space_after = Pt(7)
        if align is not None:
            p.alignment = align
    return shape


def base(title, eyebrow="NELLAI GREEN & CIVIC", subtitle="", dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(DARK if dark else CREAM)
    txt(slide, eyebrow, 0.65, 0.45, 11.9, 0.25, 10, "BDD09A" if dark else MUTED, True)
    txt(slide, title, 0.65, 0.94, 12, 0.95, 31, WHITE if dark else INK, True)
    if subtitle:
        txt(slide, subtitle, 0.65, 1.88, 12, 0.6, 14, "C3D1B6" if dark else MUTED)
    box(slide, 0.65, 7.04, 12.03, 0.012, "50704D" if dark else "DDE5D4", rounded=False)
    txt(
        slide,
        "ME PROJECT  /  LOCALHOST DEMONSTRATION  /  EDITABLE PRESENTATION",
        0.65,
        7.17,
        10.8,
        0.17,
        8,
        "A7BE8E" if dark else MUTED,
    )
    txt(
        slide,
        f"{len(prs.slides):02}",
        12,
        7.12,
        0.65,
        0.3,
        11,
        "B4C99C" if dark else MUTED,
        align=PP_ALIGN.RIGHT,
    )
    return slide


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def card(slide, title, body, x, y, w=3.86, h=2.55, tone=WHITE, number=None):
    box(slide, x, y, w, h, tone)
    if number:
        txt(slide, number, x + 0.23, y + 0.18, 0.8, 0.4, 23, ORANGE, True)
    ty = y + 0.7 if number else y + 0.25
    txt(slide, title, x + 0.23, ty, w - 0.46, 0.8, 19, GREEN, True)
    txt(slide, body, x + 0.23, ty + 0.83, w - 0.46, h - (ty - y) - 0.9, 15, MUTED)


def bullets(slide, items, x, y, w=11.9, gap=0.72, size=20):
    for i, item in enumerate(items):
        box(slide, x, y + i * gap + 0.1, 0.085, 0.085, ORANGE, rounded=False)
        txt(slide, item, x + 0.27, y + i * gap, w - 0.27, gap - 0.08, size)


def arrow(slide, x1, y1, x2, y2, stroke=MUTED):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color(stroke)
    conn.line.width = Pt(1.8)
    # Native line + editable triangular arrow head.
    if abs(x2 - x1) >= abs(y2 - y1):
        tip = slide.shapes.add_shape(
            MSO_SHAPE.CHEVRON, Inches(x2 - 0.12), Inches(y2 - 0.08), Inches(0.14), Inches(0.16)
        )
        if x2 < x1:
            tip.rotation = 180
    else:
        tip = slide.shapes.add_shape(
            MSO_SHAPE.DOWN_ARROW, Inches(x2 - 0.07), Inches(y2 - 0.12), Inches(0.14), Inches(0.16)
        )
        if y2 < y1:
            tip.rotation = 180
    tip.fill.solid()
    tip.fill.fore_color.rgb = color(stroke)
    tip.line.fill.background()


def image_fit(slide, name, x, y, w, h):
    path = OUT / "screenshots" / name
    if not path.is_file():
        box(slide, x, y, w, h, SAGE)
        txt(
            slide,
            f"Screenshot: {name}\nRun browser tests to capture the live application.",
            x + 0.3,
            y + 0.4,
            w - 0.6,
            h - 0.8,
            18,
        )
        return
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(
        str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), width=Inches(pw), height=Inches(ph)
    )


def table(slide, headers, rows, x=0.65, y=2.5, w=12.0, h=3.5, widths=None):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    t = shape.table
    if widths:
        for c, width in zip(t.columns, widths):
            c.width = Inches(width)
    for i, row in enumerate([headers] + rows):
        for j, value in enumerate(row):
            c = t.cell(i, j)
            c.text = value
            c.fill.solid()
            c.fill.fore_color.rgb = color(GREEN if i == 0 else (WHITE if i % 2 else SAGE))
            c.margin_left = c.margin_right = Inches(0.16)
            c.margin_top = Inches(0.12)
            for p in c.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(14)
                p.font.bold = i == 0
                p.font.color.rgb = color(WHITE if i == 0 else INK)
    return t


# 1 — title, entirely editable.
s = base("Nellai Green & Civic", "ME PROJECT PRESENTATION", dark=True)
txt(
    s,
    "Civic complaint management\nwith authority email alerts",
    0.7,
    2.0,
    8.3,
    1.4,
    30,
    "E4EDD5",
    True,
)
txt(s, "Report. Notify. Respond. Verify.", 0.7, 3.65, 8, 0.5, 21, "B7CB92")
box(s, 9.8, 1.65, 2.8, 3.1, "2D573C")
txt(s, "N", 10.42, 1.95, 1.6, 1.4, 93, "B7D28E", True, PP_ALIGN.CENTER)
txt(s, "OUR NELLAI\nOUR RESPONSIBILITY", 10.03, 3.68, 2.33, 0.72, 13, "E1ECD2", True, PP_ALIGN.CENTER)
txt(
    s,
    "M.E. Computer Science & Engineering\nWorking localhost demonstration\nTirunelveli / Nellai\nWeb application, email outbox and community modules",
    0.7,
    4.9,
    11.7,
    1.7,
    17,
    "DCE8CF",
)
notes(
    s,
    "Introduce yourself and your college. This localhost prototype manages complaints and sends authority alerts. Gmail accepted test NGC-063524AD on 7 October 2026; inbox arrival is checked separately. The ZIP supports a Python 3.12 setup on Windows, Linux and macOS. AI classification is a future extension.",
)

s = base(
    "A report should start a response.",
    "01 / PROBLEM DEFINITION",
    "Design problem: civic reports and environmental participation need one traceable workflow.",
)
card(
    s,
    "Scattered reporting",
    "A concern needs a location, evidence, category and a clear record.",
    0.65,
    2.65,
    number="01",
)
card(
    s,
    "Unclear responsibility",
    "Routing should follow configured teams and jurisdictions.",
    4.74,
    2.65,
    number="02",
)
card(
    s,
    "Missing closure",
    "A status change should be supported by evidence and citizen feedback.",
    8.83,
    2.65,
    number="03",
)
txt(
    s,
    "Project response: a shared, auditable platform that connects reporting with action and community care.",
    0.8,
    5.8,
    11.7,
    0.85,
    22,
    GREEN,
    True,
)
notes(
    s,
    "Describe the problem as the motivation for this design. Do not claim measured failure rates or comparisons with government systems; no such study was conducted.",
)

s = base("One core platform. Several community outcomes.", "02 / AIM & OBJECTIVES")
bullets(
    s,
    [
        "Capture civic and environmental issues with photographs and location.",
        "Assign responsibility using configurable rules and working-day deadlines.",
        "Track evidence, escalation and citizen verification.",
        "Support volunteers, tree records and water observations.",
        "Automatically queue authority email alerts for every new complaint.",
    ],
    0.85,
    2.35,
    gap=0.77,
    size=21,
)
notes(
    s,
    "Emphasize accountable workflow and extensibility. Environmental modules demonstrate how the same core can support additional use cases.",
)

s = base(
    "A focused, working localhost demonstration.",
    "03 / IMPLEMENTED SCOPE",
    "A complete complaint workflow with an administrator email outbox; no sensor hardware is required.",
)
for title, body, x, y, tone in [
    ("Citizen portal", "Register, report, track, verify and manage the account.", 0.65, 2.65, WHITE),
    ("Authority workspace", "Assigned issues, progress updates and resolution evidence.", 4.74, 2.65, WHITE),
    ("Administrator controls", "Teams, calendars, rules, modules and complaint review.", 8.83, 2.65, WHITE),
    ("GIS & environment", "Clusters, heatmap, trees and water-resource layers.", 0.65, 4.75, SAGE),
    (
        "Community participation",
        "Volunteer profiles, activities and capacity-aware joining.",
        4.74,
        4.75,
        SAGE,
    ),
    ("Authority email outbox", "Automatic alerts, message previews and delivery status.", 8.83, 4.75, SAGE),
]:
    card(s, title, body, x, y, h=1.85, tone=tone)
notes(
    s,
    "Explain that records are actually persisted and roles are enforced through the API. Seeded records are synthetic and labelled. AI and sensors are not simulated as real predictions.",
)

s = base("Open technologies, clear responsibilities.", "04 / TECHNOLOGY SELECTION")
table(
    s,
    ["Layer", "Technology", "Responsibility"],
    [
        ["Web interface", "React + TypeScript", "Role-based views, forms and reusable components"],
        ["API and services", "Python + FastAPI", "Validation, authentication and complaint workflow"],
        ["Persistence / GIS", "PostgreSQL + PostGIS", "Records, distance lookup and geographic boundaries"],
        ["Maps", "Leaflet + OpenStreetMap", "Interactive layers, clustering and heatmaps"],
        [
            "Development / tests",
            "VS Code, pytest, Playwright",
            "Local run tasks, API and browser verification",
        ],
    ],
    h=3.65,
    widths=[2.25, 3.55, 6.2],
)
txt(
    s,
    "SQLite is available as an easier local adapter; PostGIS is used for the prepared demo.",
    0.75,
    6.4,
    11.8,
    0.35,
    13,
    MUTED,
)
notes(
    s,
    "These are implementation choices, not claims of novel algorithms. Official technology references are listed at the end of this presentation.",
)

s = base("A modular monolith with API boundaries.", "05 / SYSTEM ARCHITECTURE")
for title, x in [("Citizen", 0.85), ("Authority", 4.73), ("Administrator", 8.61)]:
    box(s, x, 2.05, 3.7, 0.55, SAGE)
    txt(s, title, x, 2.15, 3.7, 0.3, 16, GREEN, True, PP_ALIGN.CENTER)
box(s, 0.85, 2.98, 11.46, 0.6, GREEN)
txt(s, "REACT WEB APPLICATION  →  REST API /api/v1", 1, 3.13, 11.1, 0.35, 17, WHITE, True, PP_ALIGN.CENTER)
box(s, 0.85, 4.03, 7.35, 0.96, WHITE)
txt(s, "CORE PLATFORM", 1.08, 4.2, 6.9, 0.3, 18, GREEN, True)
txt(s, "Auth · RBAC · Workflow · GIS · Notifications · Audit", 1.08, 4.61, 6.9, 0.25, 14, MUTED)
box(s, 8.52, 4.03, 3.79, 0.96, SAGE)
txt(s, "MODULE REGISTRY", 8.72, 4.19, 3.3, 0.3, 16, GREEN, True)
txt(s, "Green · Volunteer · Water · Flood", 8.72, 4.61, 3.32, 0.3, 12, MUTED)
box(s, 0.85, 5.63, 7.35, 0.85, GREEN)
txt(s, "POSTGRESQL + POSTGIS", 1.02, 5.87, 6.96, 0.3, 18, WHITE, True, PP_ALIGN.CENTER)
box(s, 8.52, 5.63, 3.79, 0.85, GOLD)
txt(s, "Email + escalation worker", 8.72, 5.83, 3.32, 0.45, 15, GREEN, True, PP_ALIGN.CENTER)
for x in [2.7, 6.58, 10.46]:
    arrow(s, x, 2.61, x, 2.91)
arrow(s, 4.45, 3.6, 4.45, 3.96)
arrow(s, 10.42, 3.6, 10.42, 3.96)
arrow(s, 4.45, 5.02, 4.45, 5.56)
arrow(s, 10.42, 5.6, 10.42, 5.04)
notes(
    s,
    "Every diagram element is editable. Explain that one API owns transactions, while a separate worker checks deadlines. The modules share core services and the database through documented interfaces.",
)

s = base("Access follows responsibility.", "06 / ROLES & PERMISSIONS")
table(
    s,
    ["Action", "Citizen", "Authority", "Administrator"],
    [
        ["Submit a report", "Yes", "Yes", "Yes"],
        ["View private complaint details", "Own reports", "Assigned team", "Managed reports"],
        ["Assign team / correct category", "No", "No", "Yes"],
        ["Update and resolve", "No", "Assigned team", "Yes"],
        ["Verify resolution", "Reporting citizen", "Only own reports", "Only own reports"],
        ["Configure modules / calendars", "No", "No", "Yes"],
    ],
    h=3.7,
    widths=[4.2, 2.6, 2.6, 2.6],
)
txt(
    s,
    "Staff role grants require a super administrator. Public maps omit contact details and private photos.",
    0.75,
    6.4,
    11.8,
    0.4,
    14,
    MUTED,
)
notes(
    s,
    "This matrix is enforced at API level, not just by hiding browser buttons. Resolution verification is restricted to the original reporting citizen, even when the caller is an administrator.",
)

s = base("A traceable record behind every action.", "07 / DATA MODEL")
items = [
    ("USERS", 0.8, 2.4, 2.9),
    ("COMPLAINTS", 5.2, 2.4, 3.1),
    ("AUTHORITIES", 9.7, 2.4, 2.8),
    ("UPLOADS", 0.8, 4.0, 2.9),
    ("EVENTS + ESCALATIONS", 5.2, 4.0, 3.1),
    ("RULES + CALENDARS", 9.7, 4.0, 2.8),
]
for label, x, y, w in items:
    box(s, x, y, w, 0.83, GREEN if label == "COMPLAINTS" else WHITE)
    txt(
        s,
        label,
        x + 0.1,
        y + 0.24,
        w - 0.2,
        0.32,
        16,
        WHITE if label == "COMPLAINTS" else GREEN,
        True,
        PP_ALIGN.CENTER,
    )
arrow(s, 3.74, 2.81, 5.13, 2.81)
arrow(s, 9.65, 2.81, 8.39, 2.81)
arrow(s, 2.25, 3.24, 2.25, 3.91)
arrow(s, 6.75, 3.24, 6.75, 3.91)
arrow(s, 11.1, 3.91, 11.1, 3.25)
box(s, 0.8, 5.4, 11.7, 0.94, SAGE)
txt(
    s,
    "Email outbox + audit trail  |  Optional records: trees · volunteers · water · gauge readings",
    1.1,
    5.65,
    11.1,
    0.55,
    18,
    GREEN,
    True,
    PP_ALIGN.CENTER,
)
notes(
    s,
    "Foreign keys preserve references. A dedicated complaint_emails table holds pending, sent or failed authority alerts with attempts and timestamps. Schema changes use Alembic migrations.",
)

s = base("A complaint ends with citizen verification.", "08 / COMPLAINT LIFECYCLE")
steps = [
    ("01", "Reported"),
    ("02", "Assigned"),
    ("03", "Acknowledged"),
    ("04", "In progress"),
    ("05", "Resolved"),
    ("06", "Verification"),
    ("07", "Closed"),
]
for i, (n, label) in enumerate(steps):
    x = 0.7 + i * 1.79
    box(s, x, 2.85, 1.6, 1.46, GREEN if i == 6 else WHITE)
    txt(s, n, x + 0.16, 3.04, 1.3, 0.4, 25, ORANGE if i < 6 else "CEE0AA", True)
    txt(s, label, x + 0.12, 3.7, 1.36, 0.5, 11, WHITE if i == 6 else GREEN, True)
    if i < 6:
        arrow(s, x + 1.62, 3.56, x + 1.75, 3.56)
box(s, 2.15, 5.03, 9.05, 1.05, SAGE)
txt(
    s, "Not resolved?  Citizen reopens → In progress", 2.45, 5.23, 8.45, 0.4, 22, GREEN, True, PP_ALIGN.CENTER
)
txt(
    s,
    "Resolution requires a description, photograph, timestamp, location and responsible user.",
    0.8,
    6.43,
    11.8,
    0.4,
    14,
    MUTED,
)
notes(
    s,
    "The resolve endpoint records RESOLVED and VERIFICATION events in the same transaction. Rejection is a separate administrator decision with a recorded reason. This demo does not insert a fake AI-analysis stage.",
)

s = base(
    "Deadlines follow a configured working calendar.",
    "09 / AUTHORITY ASSIGNMENT & ESCALATION",
    "Category + GPS + jurisdiction + responsibility rule → responsible team",
)
box(s, 0.7, 2.6, 7.15, 3.7, WHITE)
txt(s, "Example: 3 working days", 1, 2.9, 6.5, 0.5, 24, GREEN, True)
txt(
    s,
    "Reported: Friday, 18 September 2026\nWeekend: Saturday and Sunday\nConfigured example holiday: Monday, 21 September\nDeadline: Thursday, 24 September",
    1,
    3.66,
    6.4,
    1.9,
    18,
)
txt(
    s,
    "The holiday is a test-calendar input, not an assertion about official holidays.",
    1,
    5.69,
    6.4,
    0.47,
    12,
    MUTED,
)
card(
    s,
    "Escalation policy",
    "1. Reminder at the due check\n2. Configured higher team\n3. Additional configured levels",
    8.17,
    2.6,
    w=4.46,
    h=3.7,
    tone=SAGE,
)
notes(
    s,
    "This exact working-day calculation is covered by a test. The worker does not hard-code a government officer. Existing complaints keep a policy snapshot; reopening begins a new escalation cycle.",
)

s = base("Location informs the response.", "10 / GIS & DUPLICATE REVIEW")
card(
    s,
    "Nearby search",
    "PostGIS distance queries use metres and spatial indexes.",
    0.65,
    2.62,
    h=2.55,
    number="01",
)
card(
    s,
    "Jurisdiction matching",
    "A point is matched to configured geographic boundaries.",
    4.74,
    2.62,
    h=2.55,
    number="02",
)
card(
    s,
    "Possible duplicates",
    "Same category, within 100 m and seven days; text similarity supports review.",
    8.83,
    2.62,
    h=2.55,
    number="03",
)
txt(
    s,
    "Reports are linked after administrator review. Original reports and evidence are retained.",
    0.8,
    5.68,
    11.8,
    0.7,
    21,
    GREEN,
    True,
)
txt(
    s,
    "Public map positions are rounded. Official boundaries and ownership need verified records.",
    0.8,
    6.42,
    11.8,
    0.35,
    13,
    MUTED,
)
notes(
    s,
    "Map layers include clusters and a heatmap. The duplicate score is text similarity, not a trained image model. No report is automatically deleted as a duplicate.",
)

s = base(
    "The working dashboard.",
    "11 / LIVE APPLICATION SCREENSHOT",
    "Actual screenshot from the running local demo; all shown records are demonstration data.",
)
image_fit(s, "dashboard.png", 0.7, 2.5, 8.3, 4.28)
card(
    s,
    "A shared overview",
    "Report counts\nOpen and closed issues\nRegistered trees\nNeighbourhood map\nCategory distribution",
    9.15,
    2.5,
    w=3.48,
    h=4.28,
    tone=SAGE,
)
notes(
    s,
    "Use this screenshot if the live environment is unavailable. The figures are synthetic demo record counts, not research findings or real district statistics.",
)

s = base("Evidence and history stay with the report.", "12 / VERIFIED END-TO-END WORKFLOW")
image_fit(s, "complaint-workflow.png", 0.7, 2.12, 8.3, 4.7)
card(
    s,
    "Browser-tested sequence",
    "Citizen uploads a photo\nAuthority acknowledges\nWork is completed\nResolution photo is added\nCitizen confirms closure",
    9.15,
    2.35,
    w=3.48,
    h=4.18,
    tone=SAGE,
)
notes(
    s,
    "This is the record created by the Playwright browser test. Explain that the backend separately tests reopening, authorization failures and escalation behavior.",
)

s = base("Every new complaint creates an authority alert.", "AUTHORITY EMAIL / DESIGN", "Complaint submission stays fast because the email is delivered by a separate worker.")
for i, (title, body) in enumerate([
    ("1. Save complaint", "Report, evidence and an email queue entry are saved together."),
    ("2. Preview alert", "The administrator sees the recipient, complaint details and status."),
    ("3. Deliver or retry", "With outgoing mail configured, the worker sends and retries failures."),
]):
    card(s, title, body, 0.7 + i * 4.1, 2.65, w=3.75, h=2.55, tone=WHITE)
txt(s, "Recipient for this demo: sandymuthukutti@gmail.com", 0.9, 5.65, 11.8, 0.45, 22, GREEN, True)
txt(s, "Verified 7 October 2026: Gmail accepted test NGC-063524AD. Outbox status: SENT.", 0.9, 6.3, 11.8, 0.38, 15, MUTED)
notes(s, "A complaint creates an outbox entry even without a matched authority. Gmail accepted the labelled test email for sandymuthukutti@gmail.com and the restarted worker also delivered the other queued demo alert. SENT means SMTP acceptance, not proof of inbox arrival. Delivery needs internet and a running worker. Retries stop after five failures. The ZIP excludes the App Password, so a new PC requires private mail setup.")

s = base("The authority email, ready to review.", "AUTHORITY EMAIL / LIVE SCREENSHOT", "Actual administrator outbox: recipient, message contents and truthful delivery status.")
image_fit(s, "email-outbox.png", 0.7, 2.48, 8.65, 4.3)
card(s, "What to demonstrate", "Select the test report\nShow the recipient\nRead the email details\nExplain the SENT state\nOpen its complaint", 9.5, 2.48, w=3.1, h=4.3, tone=SAGE)
notes(s, "Sign in as administrator and open Authority email outbox. The screenshot records the real original-PC SENT state after Gmail accepted the message. Ask the recipient to check Inbox/Spam; do not treat SMTP acceptance as inbox confirmation. On another PC create a new report, because the original database is not included in the ZIP. The preview uses the actual outgoing email formatter.")

s = base("From a planting day to a living record.", "13 / GREEN NELLAI & VOLUNTEERS")
image_fit(s, "trees.png", 0.7, 2.3, 5.8, 3.6)
image_fit(s, "volunteers.png", 6.78, 2.3, 5.8, 3.6)
txt(s, "Trees: species, planting date, coordinates and care status.", 0.8, 6.11, 5.5, 0.6, 17, GREEN, True)
txt(s, "Volunteers: profiles, activities and available places.", 6.9, 6.11, 5.5, 0.6, 17, GREEN, True)
notes(
    s,
    "Survival rates use inspected trees only; newly planted trees are not counted as known survivors. No automatic species recommendation or land ownership claim is made.",
)

s = base("A module can add a new community capability.", "14 / OPEN MODULE ARCHITECTURE")
image_fit(s, "modules.png", 0.7, 2.28, 7.3, 4.2)
box(s, 8.2, 2.5, 4.42, 3.95, GREEN)
txt(s, "MODULE PACKAGE", 8.47, 2.79, 3.9, 0.4, 15, "C9DAAE", True)
txt(
    s,
    "module.json\nbackend / API / models\nfrontend component\ndatabase migration\ntests + README",
    8.47,
    3.48,
    3.9,
    2.35,
    21,
    WHITE,
)
notes(
    s,
    "Demonstrate enabling the example flood module. It adds a page and APIs for manual gauge readings. It is a reviewed installed source package, not executable code uploaded by anonymous users.",
)

s = base("Security and privacy are part of the workflow.", "15 / ENGINEERING CONTROLS")
for title, body, x, y in [
    ("Authentication", "Argon2 passwords; JWT sessions; refresh rotation.", 0.65, 2.45),
    ("Authorization", "Role and ownership checks on API operations.", 4.74, 2.45),
    ("Private evidence", "Validated images; EXIF removal; authenticated access.", 8.83, 2.45),
    ("Public-data minimization", "No public contact details or private photographs.", 0.65, 4.57),
    ("Traceability", "Workflow history, audit records and session revocation.", 4.74, 4.57),
    ("User choice", "Password-confirmed citizen account deletion.", 8.83, 4.57),
]:
    card(s, title, body, x, y, h=1.9, tone=WHITE)
notes(
    s,
    "These are implemented controls, not an independent security certification. Account deletion has documented demo limits for retained operational history. A public rollout needs further review.",
)

s = base(
    "Verification with reproducible tests.",
    "16 / VALIDATION RESULTS",
    "Results below are functional checks from this implementation, not field-performance research.",
)
box(s, 0.7, 2.6, 3.75, 2.03, GREEN)
txt(s, "25", 0.95, 2.83, 3.25, 0.9, 62, WHITE, True, PP_ALIGN.CENTER)
txt(
    s,
    "API / module tests passed\non PostgreSQL + PostGIS",
    0.95,
    3.77,
    3.25,
    0.6,
    15,
    "D3E2BC",
    align=PP_ALIGN.CENTER,
)
box(s, 4.77, 2.6, 3.75, 2.03, GREEN)
txt(s, "3", 5.02, 2.83, 3.25, 0.9, 62, WHITE, True, PP_ALIGN.CENTER)
txt(
    s,
    "Browser scenarios passed\nworkflow + outbox + layout",
    5.02,
    3.77,
    3.25,
    0.6,
    15,
    "D3E2BC",
    align=PP_ALIGN.CENTER,
)
box(s, 8.85, 2.6, 3.75, 2.03, GREEN)
txt(s, "PASS", 9.1, 2.97, 3.25, 0.7, 41, WHITE, True, PP_ALIGN.CENTER)
txt(
    s,
    "TypeScript + web build\nPython lint + migrations",
    9.1,
    3.77,
    3.25,
    0.6,
    15,
    "D3E2BC",
    align=PP_ALIGN.CENTER,
)
bullets(
    s,
    [
        "API checks cover access controls, email retries, pending alerts and private previews.",
        "Workflow tests cover reopening, duplicate linking and escalation idempotence.",
    ],
    0.85,
    5.08,
    gap=0.73,
    size=18,
)
notes(
    s,
    "Show docs/verification.md if asked. Twenty-five backend/module tests passed on PostGIS. Three browser scenarios cover the full workflow, administration/modules/responsiveness and email outbox; the outbox scenario was repeated after real Gmail acceptance. Do not claim AI accuracy, scalability numbers or government adoption.",
)

s = base("One demo package. Three operating systems.", "17 / PORTABLE DEMONSTRATION", "Extract the ZIP, install Python 3.12, then run the launcher for your PC.")
for title, body, x in [
    ("Windows", "START_WINDOWS.bat\nDouble-click to set up\nand start the demo.", 0.7),
    ("Linux", "bash START_LINUX.sh\nRun from the extracted\nproject folder.", 4.8),
    ("macOS", "bash START_MAC.command\nRun in Terminal from\nthe extracted folder.", 8.9),
]:
    card(s, title, body, x, 2.65, w=3.73, h=2.45, tone=WHITE)
txt(s, "http://localhost:8080", 0.9, 5.45, 11.5, 0.5, 26, GREEN, True, PP_ALIGN.CENTER)
txt(s, "Built website + SQLite demo data. First setup needs internet; configure Gmail privately on each PC.", 0.85, 6.25, 11.8, 0.55, 16, MUTED)
notes(
    s,
    "Follow RUN_ON_ANOTHER_PC.md. Python 3.12 is required and packages download during first setup. No Node.js or Docker is required for the packaged demo. Fresh demo logins are generated in runtime/DEMO_CREDENTIALS.md. The ZIP excludes original passwords, database and uploaded files. SQLite supports the portable workflow; the original PC uses PostGIS. Gmail and external map tiles require internet. Windows and macOS launchers are supplied but native execution was not available on the Linux build host.",
)

s = base("A clear boundary between demo and future work.", "18 / LIMITATIONS & NEXT STEPS")
table(
    s,
    ["Working in this demonstration", "Further setup or future work"],
    [
        ["Human-selected complaint categories", "Trained image and text classification models"],
        ["Configured illustrative teams and boundaries", "Verified official datasets and government APIs"],
        ["Manual water observations and gauge records", "IoT sensor feeds and validated flood predictions"],
        ["Responsive React web interface", "Completed Flutter app and store publishing"],
        ["In-app updates and Gmail SMTP alerts", "New-PC mail setup; public deployment; mobile push"],
        ["Functional local tests", "Independent security, privacy and performance reviews"],
    ],
    h=3.76,
    widths=[5.62, 6.38],
)
notes(
    s,
    "This slide is essential: explain honestly what is implemented. The user chose a simple local demo rather than a Play Store project. No fake external system or AI output has been added.",
)

s = base("Technology references & reproducibility.", "19 / REFERENCES")
references = [
    ("FastAPI security and API documentation", "https://fastapi.tiangolo.com/tutorial/security/oauth2-jwt/"),
    ("PostGIS spatial distance queries", "https://postgis.net/docs/ST_DWithin.html"),
    ("React documentation", "https://react.dev/learn"),
    (
        "Leaflet documentation / OpenStreetMap attribution",
        "https://leafletjs.com/reference.html   |   https://www.openstreetmap.org/copyright",
    ),
    ("Source, verification, and demo steps", "README.md  ·  docs/verification.md  ·  docs/demo-guide.md"),
]
for i, (title, url) in enumerate(references):
    txt(s, title, 0.83, 2.3 + i * 0.82, 11.65, 0.35, 18, GREEN, True)
    txt(s, url, 0.83, 2.69 + i * 0.82, 11.65, 0.28, 12, MUTED)
notes(
    s,
    "These are technology documentation references. A separate academic literature survey can be added by the student if required by the college; no papers or research results are fabricated in this deck.",
)

s = base("Our Nellai. Our responsibility.", "THANK YOU", dark=True)
txt(s, "One report. A clear owner.\nA verifiable resolution.", 0.85, 2.45, 11.7, 1.8, 43, "E2ECD1", True)
txt(s, "Nellai Green & Civic", 0.9, 4.67, 10, 0.55, 25, "BBD193", True)
txt(s, "Questions & live demonstration", 0.9, 5.43, 10, 0.5, 20, "B9CBA9")
notes(
    s,
    "Invite questions. Be ready to show the live workflow, module manifest, API docs, tests, or the working-day calculation.",
)

path = OUT / "Nellai_Green_Civic_Demo_Ready.pptx"
prs.save(path)
print(f"Created {len(prs.slides)} editable slides: {path}")
